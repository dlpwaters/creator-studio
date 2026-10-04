"""Local, bounded Studio downloads. Exporting never contacts a model or source URL."""

from __future__ import annotations

import html
import importlib.util
import io
import json
import math
import re
import shutil
import subprocess
import tempfile
import textwrap
import wave
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

FORMATS = {"json", "markdown", "html", "pptx", "mp4", "srt"}
STYLE_THEMES = {
    "editorial": {
        "background": "#faf9f5",
        "canvas": "#e9e8e2",
        "ink": "#172f36",
        "muted": "#50676e",
        "accent": "#9e4f2f",
        "panel": "#e9ede9",
        "line": "#ccd2cf",
        "title_bold": True,
        "radius": 16,
        "rule": 7,
    },
    "chalkboard": {
        "background": "#17352d",
        "canvas": "#102920",
        "ink": "#f3f5e8",
        "muted": "#b9ccbf",
        "accent": "#f0cb75",
        "panel": "#25483d",
        "line": "#547568",
        "title_bold": True,
        "radius": 2,
        "rule": 3,
    },
    "minimal": {
        "background": "#ffffff",
        "canvas": "#f3f4f1",
        "ink": "#202727",
        "muted": "#66716e",
        "accent": "#536f65",
        "panel": "#f0f3f0",
        "line": "#d8dfd9",
        "title_bold": False,
        "radius": 0,
        "rule": 2,
    },
}
FONT_REGULAR = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "/usr/share/fonts/liberation/LiberationSans-Regular.ttf",
)
FONT_BOLD = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/liberation/LiberationSans-Bold.ttf",
)


class ExportError(Exception):
    def __init__(self, message: str, status_code: int = 422):
        self.message = message
        self.status_code = status_code
        super().__init__(message)


@dataclass
class ExportDownload:
    filename: str
    media_type: str
    data: bytes | None = None
    path: Path | None = None
    temporary_directory: Path | None = None

    def cleanup(self) -> None:
        if self.temporary_directory:
            shutil.rmtree(self.temporary_directory, ignore_errors=True)


def export_capabilities() -> dict[str, Any]:
    return {
        "pptx_available": importlib.util.find_spec("pptx") is not None
        and importlib.util.find_spec("PIL") is not None,
        "video_available": bool(
            importlib.util.find_spec("PIL")
            and shutil.which("ffmpeg")
            and any(Path(font).is_file() for font in FONT_REGULAR)
        ),
        "video_audio": "optional_local",
        "narration_available": bool(
            shutil.which("espeak-ng") or shutil.which("espeak")
        ),
        "video_caption_note": "Captioned local video. Font coverage varies by language; review exported text.",
    }


def _text(value: Any, limit: int, label: str) -> str:
    if not isinstance(value, str) or len(value) > limit:
        raise ExportError(f"Invalid or oversized {label}.")
    # Office XML cannot contain these control characters.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)


def _payload(artifact: Any) -> dict[str, Any]:
    raw = artifact if isinstance(artifact, dict) else artifact.model_dump(mode="json")
    if not isinstance(raw, dict):
        raise ExportError("Invalid artifact.")
    data = dict(raw)
    for key, limit in (("title", 300), ("id", 100), ("language", 80), ("style", 40)):
        data[key] = _text(data.get(key, ""), limit, key)
    data["style"] = data["style"] or "editorial"
    if data["style"] not in STYLE_THEMES:
        raise ExportError("Unsupported export style.")
    if not data["title"].strip():
        raise ExportError("Artifact title cannot be empty.")
    warnings = data.get("warnings", [])
    if not isinstance(warnings, list) or len(warnings) > 200:
        raise ExportError("Invalid review warnings.")
    data["warnings"] = [_text(item, 2000, "review warning") for item in warnings]
    cards = data.get("cards")
    if not isinstance(cards, list) or not 1 <= len(cards) <= 20:
        raise ExportError("Exports require between 1 and 20 cards.")
    data["cards"] = []
    for item in cards:
        if not isinstance(item, dict):
            raise ExportError("Invalid card.")
        card = dict(item)
        for key, limit in (("title", 300), ("body", 8000), ("notes", 12000)):
            card[key] = _text(card.get(key, ""), limit, f"card {key}")
        if not card["title"].strip():
            raise ExportError("Card title cannot be empty.")
        for key in ("question", "answer"):
            card[key] = _text(card.get(key) or "", 8000, key)
        for key, count, limit in (
            ("bullets", 20, 1000),
            ("source_ids", 200, 200),
            ("options", 20, 1000),
        ):
            values = card.get(key, [])
            if not isinstance(values, list) or len(values) > count:
                raise ExportError(f"Invalid card {key}.")
            card[key] = [_text(value, limit, key) for value in values]
        duration = card.get("duration_seconds", 10)
        if (
            isinstance(duration, bool)
            or not isinstance(duration, (int, float))
            or not math.isfinite(duration)
            or not 5 <= duration <= 60
        ):
            raise ExportError("Card duration must be between 5 and 60 seconds.")
        card["duration_seconds"] = duration
        correct = card.get("correct_option")
        if correct is not None and (
            isinstance(correct, bool)
            or not isinstance(correct, int)
            or not 0 <= correct < len(card["options"])
        ):
            raise ExportError("Invalid correct option.")
        data["cards"].append(card)
    sources = data.get("sources", [])
    if not isinstance(sources, list) or len(sources) > 200:
        raise ExportError("Invalid provenance.")
    data["sources"] = []
    for source in sources:
        if not isinstance(source, dict):
            raise ExportError("Invalid provenance.")
        data["sources"].append(
            {
                "id": _text(source.get("id", ""), 200, "source ID"),
                "title": _text(source.get("title", ""), 500, "source title"),
                "url": _safe_url(source.get("url")),
            }
        )
    return data


def _safe_url(value: Any) -> str | None:
    if (
        not isinstance(value, str)
        or len(value) > 4096
        or any(ord(c) < 33 for c in value)
    ):
        return None
    try:
        parts = urlsplit(value)
        if (
            parts.scheme in {"http", "https"}
            and parts.hostname
            and not parts.username
            and not parts.password
        ):
            return value
    except ValueError:
        pass
    return None


def _provenance(data: dict, card: dict) -> list[str]:
    sources = {source["id"]: source for source in data["sources"]}
    result = []
    for source_id in card["source_ids"]:
        source = sources.get(source_id)
        result.append(
            f"{source['title']} [{source_id}]"
            + (f" — {source['url']}" if source["url"] else "")
            if source
            else source_id
        )
    return result


def _card_text(card: dict) -> list[str]:
    lines = [card["body"]] if card["body"] else []
    lines.extend(f"• {bullet}" for bullet in card["bullets"])
    if card["question"]:
        lines.append(f"Question: {card['question']}")
    lines.extend(
        f"{index + 1}. {option}" for index, option in enumerate(card["options"])
    )
    if card["answer"]:
        lines.append(f"Answer: {card['answer']}")
    if card.get("correct_option") is not None:
        lines.append(f"Correct option: {card['correct_option'] + 1}")
    return lines


def markdown(data: dict) -> str:
    lines = [
        f"# {data['title']}",
        "",
        f"Generation: {data.get('generation', 'unknown')}",
        "",
    ]
    for index, card in enumerate(data["cards"], 1):
        lines.extend([f"## {index}. {card['title']}", "", *_card_text(card), ""])
        if card["notes"]:
            lines.extend(["### Speaker notes", "", card["notes"], ""])
        lines.extend(
            [
                "### Local narration script",
                "",
                _narration_script(card, data.get("generation", "extractive")),
                "",
            ]
        )
        lines.extend(
            [
                f"Duration: {card['duration_seconds']:g} seconds",
                "",
                "Sources:",
                *[f"- {source}" for source in _provenance(data, card)],
                "",
            ]
        )
    if data.get("warnings"):
        lines.extend(["## Review notes", "", *[str(item) for item in data["warnings"]]])
    return "\n".join(lines) + "\n"


def _html_language(value: str) -> str:
    """Resolve language names and core BCP 47 locales using existing ISO data."""
    import pycountry

    value = value.strip().replace("_", "-")
    try:
        language = pycountry.languages.lookup(value)
        return getattr(language, "alpha_2", language.alpha_3).lower()
    except LookupError:
        pass
    match = re.fullmatch(
        r"([A-Za-z]{2,3})(?:-([A-Za-z]{4}))?(?:-([A-Za-z]{2}|[0-9]{3}))?", value
    )
    if not match:
        return "und"
    primary, script, region = match.groups()
    primary = primary.lower()
    field = "alpha_2" if len(primary) == 2 else "alpha_3"
    if not pycountry.languages.get(**{field: primary}):
        return "und"
    parts = [primary]
    if script:
        script = script.title()
        if not pycountry.scripts.get(alpha_4=script):
            return "und"
        parts.append(script)
    if region:
        region = region.upper()
        if not region.isdigit() and not pycountry.countries.get(alpha_2=region):
            return "und"
        parts.append(region)
    return "-".join(parts)


def html_deck(data: dict) -> str:
    esc = html.escape
    theme = STYLE_THEMES[data["style"]]
    sections = []
    for index, card in enumerate(data["cards"], 1):
        body = f'<p class="body">{esc(card["body"])}</p>' if card["body"] else ""
        bullets = "".join(f"<li>{esc(item)}</li>" for item in card["bullets"])
        question = (
            f"<p><strong>Question</strong> {esc(card['question'])}</p>"
            if card["question"]
            else ""
        )
        options = "".join(f"<li>{esc(item)}</li>" for item in card["options"])
        answer = card["answer"]
        if card.get("correct_option") is not None:
            answer += f" (Correct option: {card['correct_option'] + 1})"
        answer_html = (
            f"<details><summary>Show answer</summary><p>{esc(answer)}</p></details>"
            if answer
            else ""
        )
        sources = "".join(f"<li>{esc(item)}</li>" for item in _provenance(data, card))
        notes = (
            f'<details class="notes"><summary>Speaker notes / script</summary><p>{esc(card["notes"])}</p></details>'
            if card["notes"]
            else ""
        )
        warnings = "".join(f"<li>{esc(item)}</li>" for item in data["warnings"])
        review = f"<strong>Review notes</strong><ul>{warnings}</ul>" if warnings else ""
        sections.append(
            f'<section class="slide" aria-label="Slide {index}" tabindex="-1"><div class="eyebrow">{index:02d} / {len(data["cards"]):02d} · {esc(data.get("generation", ""))}</div><h1>{esc(card["title"])}</h1>{body}<ul>{bullets}</ul>{question}<ol>{options}</ol>{answer_html}{notes}<footer><strong>Sources</strong><ul>{sources}</ul>{review}</footer></section>'
        )
    css = """
    :root{color-scheme:light}*{box-sizing:border-box}body{margin:0;background:#e9e8e2;color:#172f36;font-family:system-ui,sans-serif}main{max-width:1200px;margin:auto;padding:32px 24px 100px}.slide{background:#faf9f5;border-top:7px solid #9e4f2f;padding:48px;min-height:640px;margin-bottom:24px;box-shadow:0 3px 18px #172f3610}.eyebrow{color:#50676e;letter-spacing:.12em;font-size:13px}h1{font-size:clamp(28px,4vw,52px);line-height:1.12;max-width:100%;overflow-wrap:anywhere}p,li{font-size:22px;line-height:1.55;overflow-wrap:anywhere}.body,p{white-space:pre-wrap}li{margin-bottom:.55em}footer{margin-top:42px;border-top:1px solid #ccd2cf;padding-top:16px}footer,footer li{font-size:13px;line-height:1.5}details{margin:20px 0}summary{cursor:pointer;font-size:16px}.notes p{font-size:17px}nav{position:fixed;bottom:0;left:0;width:100%;background:#faf9f5;border-top:1px solid #ccd2cf;padding:12px;display:flex;justify-content:center;align-items:center;gap:14px}button{border:1px solid #50676e;background:transparent;color:#172f36;border-radius:6px;padding:10px 15px;cursor:pointer;font:inherit}button:focus-visible,summary:focus-visible{outline:3px solid #9e4f2f;outline-offset:3px}body.presenting main{max-width:100%;padding:20px 20px 90px}body.presenting .slide{display:none;max-height:calc(100vh - 110px);overflow:auto;min-height:0;margin:0}body.presenting .slide.active{display:block}@media(max-width:600px){main{padding:12px 12px 100px}.slide{padding:25px;min-height:0}p,li{font-size:18px}nav{gap:7px;flex-wrap:wrap}button{padding:8px}}@media print{@page{size:landscape;margin:12mm}body{background:white}main,body.presenting main{padding:0;max-width:none}nav{display:none}.slide,body.presenting .slide{display:block;box-shadow:none;border-top:3px solid #9e4f2f;min-height:0;max-height:none;overflow:visible;break-after:page;padding:12px;margin:0}.slide:last-child{break-after:auto}h1{font-size:30px}p,li{font-size:17px}details p{display:block}.notes{display:none}footer{margin-top:24px}}
    """
    colors = {
        STYLE_THEMES["editorial"][key]: theme[key]
        for key in ("background", "canvas", "ink", "muted", "accent", "panel", "line")
    }
    colors["#172f3610"] = theme["ink"] + "10"
    css = re.sub(
        r"#[0-9a-f]{6}(?:10)?", lambda match: colors.get(match[0], match[0]), css
    )
    css += f"h1{{font-weight:{700 if theme['title_bold'] else 400}}}.slide{{border-top-width:{theme['rule']}px}}"
    if data["style"] == "minimal":
        css += ".slide{box-shadow:none}"
    # Printed dark decks need readable text even when the browser omits backgrounds.
    css += "@media print{body,.slide,body.presenting .slide{background:#fff;color:#172f36}.eyebrow{color:#50676e}}"
    script = """
    (()=>{const slides=Array.from(document.querySelectorAll('.slide'));let index=0;const status=document.getElementById('status');function show(next){index=Math.max(0,Math.min(slides.length-1,next));slides.forEach((slide,i)=>slide.classList.toggle('active',i===index));status.textContent=(index+1)+' / '+slides.length;if(document.body.classList.contains('presenting'))slides[index].focus({preventScroll:true});else slides[index].scrollIntoView({behavior:'instant',block:'start'});}document.getElementById('prev').onclick=()=>show(index-1);document.getElementById('next').onclick=()=>show(index+1);document.getElementById('present').onclick=()=>{document.body.classList.toggle('presenting');show(index);};document.getElementById('print').onclick=()=>window.print();document.addEventListener('keydown',event=>{if(event.target.closest('button,summary,input,textarea,select'))return;if(['ArrowRight','PageDown','ArrowLeft','PageUp','Home','End','Escape'].includes(event.key)){event.preventDefault();if(event.key==='Escape'){document.body.classList.remove('presenting');return;}show(event.key==='Home'?0:event.key==='End'?slides.length-1:['ArrowRight','PageDown'].includes(event.key)?index+1:index-1);}});show(0);})();
    """
    return f'<!doctype html><html lang="{_html_language(data["language"])}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{esc(data["title"])}</title><style>{css}</style></head><body data-style="{data["style"]}"><main>{"".join(sections)}</main><nav aria-label="Presentation controls"><button id="prev">Previous</button><span id="status" aria-live="polite"></span><button id="next">Next</button><button id="present">Present</button><button id="print">Print / PDF</button></nav><script>{script}</script></body></html>'


def _time(seconds: float) -> str:
    milliseconds = round(seconds * 1000)
    hours, milliseconds = divmod(milliseconds, 3600000)
    minutes, milliseconds = divmod(milliseconds, 60000)
    whole, milliseconds = divmod(milliseconds, 1000)
    return f"{hours:02}:{minutes:02}:{whole:02},{milliseconds:03}"


def _narration_script(card: dict, generation: str = "ai") -> str:
    if generation == "ai" and card["notes"]:
        return card["notes"]
    return "\n".join([card["title"], *_card_text(card)])


def caption_cues(
    data: dict, narration: bool = False
) -> list[tuple[float, float, int, str]]:
    """Match burned-in captions and SRT. Long cards disclose abbreviated visuals."""
    cues = []
    start = 0.0
    for card_index, card in enumerate(data["cards"]):
        paragraphs = (
            [_narration_script(card, data.get("generation", "extractive"))]
            if narration
            else [card["title"], *_card_text(card)]
        )
        parts = []
        for paragraph in paragraphs:
            parts.extend(
                textwrap.wrap(
                    " ".join(paragraph.split()),
                    width=155,
                    break_long_words=True,
                    break_on_hyphens=False,
                )
            )
        maximum = max(2, int(card["duration_seconds"] // 2.5))
        if len(parts) > maximum and not narration:
            parts = parts[: maximum - 1] + [
                "Visual excerpt abbreviated. Read the complete card and script in the Markdown export."
            ]
        card_start = start
        spoken_duration = (
            card.get("_speech_seconds", card["duration_seconds"])
            if narration
            else card["duration_seconds"]
        )
        for part_index, part in enumerate(parts):
            end = card_start + spoken_duration * (part_index + 1) / len(parts)
            if part_index == len(parts) - 1:
                # Hold the last caption during padded silence; earlier cues follow speech.
                end = card_start + card["duration_seconds"]
            cues.append((start, end, card_index, part))
            start = end
    return cues


def srt(data: dict, narration: bool = False) -> str:
    return (
        "\n\n".join(
            f"{index}\n{_time(start)} --> {_time(end)}\n{text}"
            for index, (start, end, _, text) in enumerate(
                caption_cues(data, narration), 1
            )
        )
        + "\n"
    )


def pptx(data: dict) -> bytes:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise ExportError(
            "PowerPoint export requires python-pptx in the backend.", 503
        ) from exc
    presentation = Presentation()
    theme = STYLE_THEMES[data["style"]]
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    dark, muted, accent = (
        RGBColor.from_string(theme["ink"].lstrip("#")),
        RGBColor.from_string(theme["muted"].lstrip("#")),
        RGBColor.from_string(theme["accent"].lstrip("#")),
    )

    def box(slide, text, x, y, width, height, size, color=dark, bold=False):
        shape = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(width), Inches(height)
        )
        frame = shape.text_frame
        frame.word_wrap = True
        for index, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "DejaVu Sans"
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = color
            paragraph.font.bold = bold
            paragraph.space_after = Pt(5)
            paragraph.line_spacing = Pt(size * 1.2)
        return shape

    for index, card in enumerate(data["cards"], 1):
        # Manually wrap and paginate; no rasterized text or hidden overflow.
        title_lines = _wrap_pixels(
            card["title"], width=820, size=28, bold=theme["title_bold"]
        ) or ["Untitled"]
        title_pages = [
            title_lines[offset : offset + 3] for offset in range(0, len(title_lines), 3)
        ]
        lines = []
        for paragraph in _card_text(card):
            for line in paragraph.splitlines() or [""]:
                lines.extend(_wrap_pixels(line, width=800, size=22))
        pages = [lines[offset : offset + 8] for offset in range(0, len(lines), 8)] or [
            []
        ]
        for page_index, page in enumerate(pages):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(
                theme["background"].lstrip("#")
            )
            box(
                slide,
                f"{index:02d} / {len(data['cards']):02d}    {data.get('generation', '').upper()}",
                0.65,
                0.4,
                12,
                0.4,
                13,
                muted,
            )
            title = "\n".join(title_lines[:3])
            if len(title_pages) > 1:
                title = "\n".join(title_lines[:2]) + "… (full title in notes)"
            box(slide, title, 0.65, 0.95, 12, 1.65, 28, dark, theme["title_bold"])
            box(slide, "\n".join(page), 0.8, 2.6, 11.7, 3.8, 22)
            source_titles = [
                source.split(" — ")[0] for source in _provenance(data, card)
            ]
            footer = "Sources: " + ("; ".join(source_titles) or "No card citations")
            if data["warnings"]:
                footer += " · Review warnings in speaker notes."
            box(
                slide,
                textwrap.shorten(footer, width=155, placeholder="… (see notes)"),
                0.8,
                6.65,
                11.7,
                0.5,
                10,
                muted,
            )
            box(
                slide,
                f"{page_index + 1}/{len(pages)}",
                12.05,
                6.65,
                0.6,
                0.5,
                10,
                accent,
            )
            notes = [
                card["title"],
                "",
                "Speaker notes:",
                card["notes"],
                "",
                "Local narration script:",
                _narration_script(card, data.get("generation", "extractive")),
                "",
                "Full card:",
                *_card_text(card),
                "",
                "Source provenance:",
                *_provenance(data, card),
                "",
                f"Duration: {card['duration_seconds']:g} seconds",
                "",
                "Review warnings:",
                *data["warnings"],
            ]
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _font(size: int, bold: bool = False):
    from PIL import ImageFont

    for candidate in FONT_BOLD + FONT_REGULAR if bold else FONT_REGULAR:
        if Path(candidate).is_file():
            return ImageFont.truetype(candidate, size)
    raise ExportError("Video export needs a local TrueType font.", 503)


def _wrap_pixels(text: str, width: int, size: int, bold: bool = False) -> list[str]:
    """Use actual font metrics so wide letters and long words stay in bounds."""
    try:
        font = _font(size, bold)
    except (ImportError, ExportError):
        # Editable Office text still works without local raster-rendering fonts.
        return textwrap.wrap(text, width=max(1, int(width / size))) or [""]
    lines = []
    line = ""
    for word in text.split():
        candidate = f"{line} {word}" if line else word
        if font.getlength(candidate) <= width:
            line = candidate
            continue
        if line:
            lines.append(line)
            line = ""
        for character in word:
            if line and font.getlength(line + character) > width:
                lines.append(line)
                line = ""
            line += character
    if line:
        lines.append(line)
    return lines or [""]


def _prepare_narration(data: dict, directory: Path) -> Path:
    """Local speech only. Preserve every word; extend scene duration within bounds."""
    executable = shutil.which("espeak-ng") or shutil.which("espeak")
    if not executable:
        raise ExportError(
            "Local narration requires espeak-ng or espeak in the backend.", 503
        )
    # Explicit voices prevent shell arguments or unreviewed voice-file paths.
    voices = {
        "en": "en",
        "es": "es",
        "fr": "fr",
        "de": "de",
        "it": "it",
        "pt": "pt",
        "nl": "nl",
        "pl": "pl",
    }
    language = data["language"].lower().replace("_", "-").split("-")[0]
    language = {
        "english": "en",
        "spanish": "es",
        "french": "fr",
        "german": "de",
        "italian": "it",
        "portuguese": "pt",
        "dutch": "nl",
        "polish": "pl",
    }.get(language, language)
    if language not in voices:
        raise ExportError(
            "Local narration supports en, es, fr, de, it, pt, nl and pl. Use captions for other languages."
        )
    audio = []
    audio_format = None
    script_field = (
        "speaker notes" if data.get("generation") == "ai" else "body and bullets"
    )
    for index, card in enumerate(data["cards"]):
        script = _narration_script(card, data.get("generation", "extractive"))
        if len(script) > 2000:
            raise ExportError(
                f"A narration script is too long for a 60-second scene. Shorten the card {script_field}."
            )
        path = directory / f"speech-{index:04d}.wav"
        try:
            result = subprocess.run(
                [
                    executable,
                    "-v",
                    voices[language],
                    "-s",
                    "160",
                    "-w",
                    str(path),
                    "--stdin",
                ],
                input=script.encode("utf-8"),
                capture_output=True,
                timeout=30,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise ExportError(
                "Local speech synthesis exceeded its time limit.", 503
            ) from exc
        if result.returncode or not path.is_file():
            raise ExportError(
                "Local speech synthesis failed. Check the installed espeak voices.", 503
            )
        with wave.open(str(path), "rb") as speech:
            current_format = (
                speech.getnchannels(),
                speech.getsampwidth(),
                speech.getframerate(),
            )
            if audio_format is not None and current_format != audio_format:
                raise ExportError(
                    "Local speech returned incompatible audio formats.", 503
                )
            audio_format = current_format
            seconds = speech.getnframes() / speech.getframerate()
            if seconds + 0.4 > 60:
                raise ExportError(
                    f"A narration script exceeds the 60-second scene limit. Shorten the card {script_field}."
                )
            card["_speech_seconds"] = seconds
            card["duration_seconds"] = max(
                card["duration_seconds"], round(seconds + 0.4, 3)
            )
            frames = speech.readframes(speech.getnframes())
            silence = (
                round((card["duration_seconds"] - seconds) * speech.getframerate())
                * speech.getnchannels()
                * speech.getsampwidth()
            )
            audio.append(frames + bytes(silence))
    output = directory / "narration.wav"
    with wave.open(str(output), "wb") as combined:
        combined.setnchannels(audio_format[0])
        combined.setsampwidth(audio_format[1])
        combined.setframerate(audio_format[2])
        for frames in audio:
            combined.writeframes(frames)
    return output


def _video(data: dict, filename: str, narration: bool = False) -> ExportDownload:
    theme = STYLE_THEMES[data["style"]]
    if not export_capabilities()["video_available"]:
        raise ExportError(
            "Video export requires Pillow, ffmpeg and a local TrueType font.", 503
        )
    from PIL import Image, ImageDraw

    directory = Path(tempfile.mkdtemp(prefix="notebook-studio-"))
    download = ExportDownload(
        filename,
        "video/mp4",
        path=directory / "video.mp4",
        temporary_directory=directory,
    )
    try:
        audio_path = _prepare_narration(data, directory) if narration else None
        cues = caption_cues(data, narration)
        total = sum(card["duration_seconds"] for card in data["cards"])
        segments = []
        for index, (start, end, card_index, caption) in enumerate(cues):
            card = data["cards"][card_index]
            image = Image.new("RGB", (1280, 720), theme["background"])
            draw = ImageDraw.Draw(image)
            draw.rectangle((0, 0, 1280, 12), fill=theme["accent"])
            draw.text(
                (64, 45),
                f"{card_index + 1:02d} / {len(data['cards']):02d}    OPEN NOTEBOOK STUDIO",
                font=_font(20),
                fill=theme["muted"],
            )
            title_lines = _wrap_pixels(
                card["title"], width=1140, size=43, bold=theme["title_bold"]
            )
            title = title_lines[:3]
            if len(title_lines) > 3:
                title[-1] = title[-1][:35] + "… (full title in script)"
            draw.multiline_text(
                (64, 105),
                "\n".join(title),
                font=_font(43, theme["title_bold"]),
                fill=theme["ink"],
                spacing=12,
            )
            # Progressive cue changes reveal the explanation one thought at a time.
            draw.rounded_rectangle(
                (64, 325, 1216, 565), radius=theme["radius"], fill=theme["panel"]
            )
            lines = _wrap_pixels(caption, width=1090, size=32)
            draw.multiline_text(
                (92, 351),
                "\n".join(lines),
                font=_font(32),
                fill=theme["ink"],
                spacing=12,
            )
            source_label = "; ".join(
                source.split(" — ")[0] for source in _provenance(data, card)
            )
            draw.text(
                (64, 604),
                textwrap.shorten(
                    "Sources: " + (source_label or "See artifact provenance"),
                    width=105,
                    placeholder="…",
                ),
                font=_font(18),
                fill=theme["muted"],
            )
            label = (
                "Captioned visual explainer · local synthesized narration"
                if narration
                else "Captioned visual explainer · no narration · full script in Markdown export"
            )
            draw.text((64, 649), label, font=_font(17), fill=theme["muted"])
            draw.rectangle((0, 710, 1280, 720), fill=theme["line"])
            draw.rectangle(
                (0, 710, round(1280 * end / total), 720), fill=theme["accent"]
            )
            image_path = directory / f"frame-{index:04d}.png"
            image.save(image_path)
            segments.extend(
                [f"file '{image_path.name}'", f"duration {end - start:.6f}"]
            )
        segments.append(f"file 'frame-{len(cues) - 1:04d}.png'")
        manifest = directory / "frames.txt"
        manifest.write_text("\n".join(segments) + "\n", encoding="utf-8")
        command = [
            shutil.which("ffmpeg"),
            "-hide_banner",
            "-loglevel",
            "error",
            "-nostdin",
            "-y",
            "-f",
            "concat",
            "-safe",
            "1",
            "-i",
            str(manifest),
        ]
        if audio_path:
            command.extend(["-i", str(audio_path), "-c:a", "aac", "-b:a", "96k"])
        else:
            command.append("-an")
        command.extend(
            [
                "-vf",
                "fps=15,format=yuv420p",
                "-c:v",
                "libx264",
                "-preset",
                "veryfast",
                "-crf",
                "24",
                "-threads",
                "2",
                "-t",
                str(total),
                "-movflags",
                "+faststart",
                str(download.path),
            ]
        )
        result = subprocess.run(
            command, capture_output=True, timeout=min(420, 60 + total / 2), check=False
        )
        if (
            result.returncode
            or not download.path.is_file()
            or download.path.stat().st_size == 0
        ):
            # ffmpeg diagnostics can echo input; keep artifact content out of errors/logs.
            raise ExportError(
                "Local video rendering failed. Verify the ffmpeg H.264 encoder is available.",
                503,
            )
        for path in directory.iterdir():
            if path != download.path:
                path.unlink()
        return download
    except subprocess.TimeoutExpired as exc:
        download.cleanup()
        raise ExportError(
            "Local video rendering exceeded its time limit. Try fewer or shorter cards.",
            503,
        ) from exc
    except BaseException:
        download.cleanup()
        raise


def export_artifact(
    artifact: Any, format: str, narration: str = "none"
) -> ExportDownload:
    if format not in FORMATS:
        raise ExportError("Unsupported export format.", 404)
    data = _payload(artifact)
    if narration not in {"none", "local"}:
        raise ExportError("Unsupported narration mode.")
    stem = (
        re.sub(r"[^a-zA-Z0-9_-]+", "-", data["title"]).strip("-")[:70]
        or "studio-artifact"
    )
    extension = {"markdown": "md"}.get(format, format)
    filename = f"{stem}.{extension}"
    if format == "mp4":
        return _video(data, filename, narration == "local")
    if format == "pptx":
        return ExportDownload(
            filename,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            data=pptx(data),
        )
    if format == "json":
        content, media_type = (
            json.dumps(data, ensure_ascii=False, indent=2),
            "application/json",
        )
    elif format == "markdown":
        content, media_type = markdown(data), "text/markdown"
    elif format == "html":
        content, media_type = html_deck(data), "text/html"
    else:
        if narration == "local":
            with tempfile.TemporaryDirectory(
                prefix="notebook-studio-captions-"
            ) as directory:
                _prepare_narration(data, Path(directory))
                content = srt(data, True)
        else:
            content = srt(data)
        media_type = "application/x-subrip"
    return ExportDownload(filename, media_type, data=content.encode("utf-8"))
