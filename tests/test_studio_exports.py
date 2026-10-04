"""Real file checks for the local Studio export pipeline, without provider calls."""

import copy
import io
import json
import shutil
import subprocess
import threading
import wave
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest

from api import studio_exports as exports


@pytest.fixture
def artifact():
    return {
        "id": "a" * 32,
        "notebook_id": "notebook:test",
        "title": "Research & evidence",
        "kind": "slides",
        "language": "English",
        "style": "editorial",
        "generation": "extractive",
        "cards": [
            {
                "id": "card-1",
                "title": "Start with the evidence",
                "body": "A source provides facts. A note records your interpretation.",
                "bullets": ["Keep provenance", "Review claims"],
                "notes": "Start with the evidence. Keep provenance and review claims.",
                "source_ids": ["source:paper", "note:review"],
                "duration_seconds": 5,
            }
        ],
        "sources": [
            {
                "id": "source:paper",
                "title": "Original paper",
                "url": "https://example.org/paper",
            },
            {"id": "note:review", "title": "Review note", "url": None},
        ],
        "warnings": [],
    }


def test_json_markdown_keep_provenance_and_script(artifact):
    decoded = json.loads(exports.export_artifact(artifact, "json").data)
    assert decoded["cards"][0]["source_ids"] == ["source:paper", "note:review"]
    document = exports.export_artifact(artifact, "markdown").data.decode()
    assert "https://example.org/paper" in document
    assert "Review note [note:review]" in document
    assert "Speaker notes" in document and "Local narration script" in document
    assert artifact["cards"][0]["notes"] in document


def test_json_export_stays_valid_against_strict_saved_artifact_schema(artifact):
    from api.studio_models import StudioArtifact

    artifact.update(
        audience="Reader",
        created_at="2026-10-04T12:00:00Z",
        updated_at="2026-10-04T12:00:00Z",
        source_ids=["source:paper"],
        note_ids=["note:review"],
    )
    model = StudioArtifact.model_validate(artifact)
    payload = exports.export_artifact(model, "json").data
    exported = StudioArtifact.model_validate_json(payload)
    assert exported.cards[0].duration_seconds == 5
    assert exported.source_ids == model.source_ids


def test_html_escapes_every_user_field_and_has_offline_controls(artifact):
    attack = '</script><img src=x onerror="alert(1)">'
    artifact["title"] = attack
    artifact["language"] = 'en" onload="alert(1)'
    artifact["cards"][0].update(
        title=attack,
        body=attack,
        notes=attack,
        bullets=[attack],
        question=attack,
        answer=attack,
        options=[attack],
    )
    artifact["sources"][0]["title"] = attack
    artifact["sources"][0]["url"] = "javascript:alert(1)"
    output = exports.export_artifact(artifact, "html")
    content = output.data.decode()
    assert attack not in content
    assert "&lt;img src=x onerror=&quot;alert(1)&quot;&gt;" in content
    assert "javascript:alert(1)" not in content
    assert 'lang="und"' in content
    assert "ArrowRight" in content and "window.print()" in content
    assert "@media print" in content
    assert "<script src=" not in content
    assert "/" not in output.filename and '"' not in output.filename


def test_caption_cues_are_contiguous_bounded_and_disclose_abbreviation(artifact):
    artifact["cards"].append(copy.deepcopy(artifact["cards"][0]))
    artifact["cards"][0]["body"] = "Long content. " * 200
    data = exports._payload(artifact)
    cues = exports.caption_cues(data)
    assert cues[0][0] == 0
    assert cues[-1][1] == pytest.approx(10)
    assert all(
        first[1] == pytest.approx(second[0]) for first, second in zip(cues, cues[1:])
    )
    assert any("abbreviated" in cue[3] for cue in cues)
    subtitles = exports.export_artifact(artifact, "srt").data.decode()
    assert "00:00:00,000 -->" in subtitles
    assert "00:00:10,000" in subtitles
    assert exports._time(3661.9996) == "01:01:02,000"


@pytest.mark.parametrize(
    "change",
    [
        {"duration_seconds": float("nan")},
        {"duration_seconds": 61},
        {"duration_seconds": True},
        {"bullets": "bad"},
        {"body": "x" * 8001},
        {"correct_option": 2, "options": ["only one"]},
    ],
)
def test_malformed_card_rejected_before_render(artifact, change):
    artifact["cards"][0].update(change)
    with pytest.raises(exports.ExportError):
        exports.export_artifact(artifact, "html")


def test_bounds_and_download_names(artifact):
    artifact["title"] = '../../etc/passwd\r\n"'
    assert exports.export_artifact(artifact, "json").filename == "etc-passwd.json"
    artifact["cards"] *= 21
    with pytest.raises(exports.ExportError, match="20 cards"):
        exports.export_artifact(artifact, "json")
    with pytest.raises(exports.ExportError, match="Unsupported"):
        exports.export_artifact(artifact, "../../file")


def test_pptx_is_editable_paginated_and_has_notes_provenance(artifact):
    pytest.importorskip("pptx")
    from pptx import Presentation

    artifact["cards"][0]["body"] = (
        "A readable long sentence with source evidence. " * 60
    )
    artifact["cards"][0]["notes"] += "\x00 Invalid XML controls are removed."
    download = exports.export_artifact(artifact, "pptx")
    with zipfile.ZipFile(io.BytesIO(download.data)) as package:
        assert package.testzip() is None
        notes = [
            name
            for name in package.namelist()
            if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")
        ]
        assert notes
        all_notes = " ".join(package.read(name).decode() for name in notes)
        assert "https://example.org/paper" in all_notes
        assert "note:review" in all_notes
        for name in notes:
            ElementTree.fromstring(package.read(name))
    presentation = Presentation(io.BytesIO(download.data))
    assert len(presentation.slides) > 1
    text_shapes = [
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert text_shapes and any(
        "Start with the evidence" in shape.text for shape in text_shapes
    )
    assert not any(
        shape.shape_type == 13
        for slide in presentation.slides
        for shape in slide.shapes
    )


def test_real_small_mp4_has_h264_and_cleanup(artifact):
    if not exports.export_capabilities()["video_available"] or not shutil.which(
        "ffprobe"
    ):
        pytest.skip("Local video dependencies unavailable")
    output = exports.export_artifact(artifact, "mp4")
    directory = output.temporary_directory
    try:
        probe = subprocess.run(
            [
                shutil.which("ffprobe"),
                "-v",
                "error",
                "-show_streams",
                "-show_format",
                "-of",
                "json",
                str(output.path),
            ],
            capture_output=True,
            check=True,
        )
        metadata = json.loads(probe.stdout)
        assert metadata["streams"][0]["codec_name"] == "h264"
        assert float(metadata["format"]["duration"]) == pytest.approx(5, abs=0.15)
        assert not any(
            stream["codec_type"] == "audio" for stream in metadata["streams"]
        )
        assert sorted(path.name for path in directory.iterdir()) == ["video.mp4"]
    finally:
        output.cleanup()
    assert not directory.exists()


def test_failed_video_render_cleans_temp_and_hides_diagnostics(
    artifact, monkeypatch, tmp_path
):
    monkeypatch.setattr(exports.tempfile, "mkdtemp", lambda **kwargs: str(tmp_path))
    monkeypatch.setattr(
        exports.subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess(
            args[0], 1, stderr=b"private artifact content"
        ),
    )
    with pytest.raises(
        exports.ExportError, match="Local video rendering failed"
    ) as error:
        exports.export_artifact(artifact, "mp4")
    assert "private" not in str(error.value)
    assert not tmp_path.exists()


def test_narration_keeps_words_pads_audio_and_does_not_mutate_artifact(
    artifact, monkeypatch, tmp_path
):
    artifact["generation"] = "ai"
    calls = []

    def synthesize(command, **kwargs):
        calls.append((command, kwargs))
        target = Path(command[command.index("-w") + 1])
        with wave.open(str(target), "wb") as output:
            output.setnchannels(1)
            output.setsampwidth(2)
            output.setframerate(22050)
            output.writeframes(bytes(22050 * 2 * 7))
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(exports.shutil, "which", lambda command: "/usr/bin/espeak-ng")
    monkeypatch.setattr(exports.subprocess, "run", synthesize)
    data = exports._payload(artifact)
    audio = exports._prepare_narration(data, tmp_path)
    assert calls[0][1]["input"].decode() == artifact["cards"][0]["notes"]
    assert artifact["cards"][0]["notes"] not in calls[0][0]
    assert calls[0][0][calls[0][0].index("-v") + 1] == "en"
    assert data["cards"][0]["duration_seconds"] == 7.4
    assert artifact["cards"][0]["duration_seconds"] == 5
    with wave.open(str(audio), "rb") as output:
        assert output.getnframes() / output.getframerate() == pytest.approx(7.4)
    subtitles = exports.srt(data, True)
    assert "Keep provenance and review claims." in subtitles
    assert "00:00:07,400" in subtitles


def test_narration_missing_or_unsupported_returns_clear_error(
    artifact, monkeypatch, tmp_path
):
    monkeypatch.setattr(exports.shutil, "which", lambda command: None)
    with pytest.raises(exports.ExportError, match="requires espeak"):
        exports._prepare_narration(exports._payload(artifact), tmp_path)
    monkeypatch.setattr(exports.shutil, "which", lambda command: "/usr/bin/espeak-ng")
    artifact["language"] = "../../evil"
    with pytest.raises(exports.ExportError, match="supports en"):
        exports._prepare_narration(exports._payload(artifact), tmp_path)


def test_real_narrated_mp4_and_srt_have_same_duration(artifact):
    artifact["generation"] = "ai"
    capabilities = exports.export_capabilities()
    if (
        not capabilities["video_available"]
        or not capabilities["narration_available"]
        or not shutil.which("ffprobe")
    ):
        pytest.skip("Local narration dependencies unavailable")
    output = exports.export_artifact(artifact, "mp4", "local")
    try:
        metadata = json.loads(
            subprocess.run(
                [
                    shutil.which("ffprobe"),
                    "-v",
                    "error",
                    "-show_streams",
                    "-show_format",
                    "-of",
                    "json",
                    str(output.path),
                ],
                capture_output=True,
                check=True,
            ).stdout
        )
        assert any(stream["codec_type"] == "audio" for stream in metadata["streams"])
        subtitles = exports.export_artifact(artifact, "srt", "local").data.decode()
        assert "Keep provenance and review claims." in subtitles
        assert float(metadata["format"]["duration"]) == pytest.approx(5, abs=0.15)
    finally:
        output.cleanup()


def test_extractive_narration_reads_subject_matter_instead_of_boilerplate(artifact):
    data = exports._payload(artifact)
    data["cards"][0]["notes"] = (
        "Verbatim source excerpt. Review it in its original context."
    )
    script = exports._narration_script(data["cards"][0], "extractive")
    assert data["cards"][0]["body"] in script
    assert "Keep provenance" in script
    assert "Verbatim source excerpt" not in script
    captions = exports.srt(data, True)
    assert "A source provides facts." in captions


def test_caption_timing_follows_speech_before_padded_silence(artifact):
    artifact["generation"] = "ai"
    data = exports._payload(artifact)
    data["cards"][0]["notes"] = "One thoughtful sentence. " * 20
    data["cards"][0]["duration_seconds"] = 60
    data["cards"][0]["_speech_seconds"] = 20
    cues = exports.caption_cues(data, True)
    assert len(cues) > 1
    assert cues[-2][1] < 20
    assert cues[-1][0] < 20
    assert cues[-1][1] == 60


def test_review_warnings_retained_in_html_and_pptx_notes(artifact):
    artifact["warnings"] = ["Source input was truncated. Review its original context."]
    warning = artifact["warnings"][0]
    assert warning in exports.export_artifact(artifact, "html").data.decode()
    with zipfile.ZipFile(
        io.BytesIO(exports.export_artifact(artifact, "pptx").data)
    ) as package:
        notes = [
            name
            for name in package.namelist()
            if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")
        ]
        assert any(warning in package.read(name).decode() for name in notes)


def test_editable_pptx_can_fall_back_without_local_fonts(artifact, monkeypatch):
    monkeypatch.setattr(exports, "FONT_REGULAR", ())
    monkeypatch.setattr(exports, "FONT_BOLD", ())
    output = exports.export_artifact(artifact, "pptx")
    assert output.data.startswith(b"PK")


def test_video_bold_uses_regular_fallback_and_pptx_requires_pillow(monkeypatch):
    monkeypatch.setattr(exports, "FONT_BOLD", ())
    assert exports._font(22, bold=True).size == 22
    original = exports.importlib.util.find_spec
    monkeypatch.setattr(
        exports.importlib.util,
        "find_spec",
        lambda module: None if module == "PIL" else original(module),
    )
    capabilities = exports.export_capabilities()
    assert not capabilities["pptx_available"]
    assert not capabilities["video_available"]


@pytest.mark.parametrize("style", ["editorial", "chalkboard", "minimal"])
def test_selected_style_survives_html_pptx_and_encoded_video(artifact, style):
    from PIL import Image
    from pptx import Presentation

    artifact["style"] = style
    theme = exports.STYLE_THEMES[style]
    document = exports.export_artifact(artifact, "html").data.decode()
    assert f'data-style="{style}"' in document
    assert f"background:{theme['background']}" in document
    assert f"color:{theme['ink']}" in document
    assert f"h1{{font-weight:{700 if theme['title_bold'] else 400}}}" in document
    deck = Presentation(io.BytesIO(exports.export_artifact(artifact, "pptx").data))
    assert (
        str(deck.slides[0].background.fill.fore_color.rgb)
        == theme["background"].lstrip("#").upper()
    )
    heading = next(
        shape
        for shape in deck.slides[0].shapes
        if shape.has_text_frame and artifact["cards"][0]["title"] in shape.text
    )
    assert heading.text_frame.paragraphs[0].font.bold == theme["title_bold"]
    assert (
        str(heading.text_frame.paragraphs[0].font.color.rgb)
        == theme["ink"].lstrip("#").upper()
    )
    if not exports.export_capabilities()["video_available"]:
        pytest.skip("Local video dependencies unavailable")
    video = exports.export_artifact(artifact, "mp4")
    try:
        decoded = subprocess.run(
            [
                shutil.which("ffmpeg"),
                "-hide_banner",
                "-loglevel",
                "error",
                "-i",
                str(video.path),
                "-frames:v",
                "1",
                "-f",
                "image2pipe",
                "-vcodec",
                "png",
                "pipe:1",
            ],
            capture_output=True,
            check=True,
        )
        frame = Image.open(io.BytesIO(decoded.stdout)).convert("RGB")
        actual = frame.getpixel((20, 100))
        expected = tuple(
            int(theme["background"][offset : offset + 2], 16) for offset in (1, 3, 5)
        )
        assert all(abs(a - b) <= 5 for a, b in zip(actual, expected))
    finally:
        video.cleanup()


def test_unknown_export_style_rejected(artifact):
    artifact["style"] = "<script>alert(1)</script>"
    with pytest.raises(exports.ExportError, match="Unsupported export style"):
        exports.export_artifact(artifact, "html")


def _contrast_ratio(foreground, background):
    def luminance(color):
        components = [int(color[offset : offset + 2], 16) / 255 for offset in (1, 3, 5)]
        linear = [
            channel / 12.92
            if channel <= 0.04045
            else ((channel + 0.055) / 1.055) ** 2.4
            for channel in components
        ]
        return sum(
            channel * weight
            for channel, weight in zip(linear, (0.2126, 0.7152, 0.0722))
        )

    dark, light = sorted((luminance(foreground), luminance(background)))
    return (light + 0.05) / (dark + 0.05)


@pytest.mark.parametrize("style", ["editorial", "chalkboard", "minimal"])
@pytest.mark.parametrize("text_role", ["ink", "muted", "accent"])
def test_palette_normal_text_has_wcag_aa_contrast_on_export_surfaces(
    artifact, style, text_role
):
    theme = exports.STYLE_THEMES[style]
    # Muted colors label slides and sources; accent colors label PPTX page counts.
    for surface in ("background", "canvas", "panel"):
        assert _contrast_ratio(theme[text_role], theme[surface]) >= 4.5, (
            style,
            text_role,
            surface,
        )
    artifact["style"] = style
    document = exports.export_artifact(artifact, "html").data.decode()
    assert f".eyebrow{{color:{theme['muted']};" in document
    assert "#627980" not in document and "#bd6b45" not in document


@pytest.mark.parametrize(
    "language, expected",
    [
        ("English", "en"),
        ("French", "fr"),
        ("en-US", "en-US"),
        ("pt_br", "pt-BR"),
        ("zh-Hant-TW", "zh-Hant-TW"),
        ('en" onload="alert(1)', "und"),
        ("Unrecognized language", "und"),
    ],
)
def test_html_language_uses_valid_tags_and_unknown_fallback(
    artifact, language, expected
):
    artifact["language"] = language
    document = exports.export_artifact(artifact, "html").data.decode()
    assert f'<html lang="{expected}">' in document


@pytest.mark.asyncio
async def test_router_renders_in_worker_and_attaches_private_download_headers(
    artifact, monkeypatch
):
    from api.routers import studio_exports as routes

    async def load(artifact_id):
        assert artifact_id == artifact["id"]
        return artifact

    loop_thread = threading.get_ident()
    render = exports.export_artifact

    def check_worker(*args):
        assert threading.get_ident() != loop_thread
        return render(*args)

    monkeypatch.setattr(routes, "get_artifact", load)
    monkeypatch.setattr(routes, "export_artifact", check_worker)
    response = await routes.download_artifact(artifact["id"], "html")
    assert response.headers["cache-control"] == "private, no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-disposition"].startswith("attachment;")
    assert b"Presentation controls" in response.body


@pytest.mark.asyncio
async def test_router_maps_artifact_scoping_errors_without_rendering(monkeypatch):
    from fastapi import HTTPException

    from api.routers import studio_exports as routes
    from api.studio_service import StudioError

    async def inaccessible(artifact_id):
        raise StudioError(404, "Artifact not found")

    monkeypatch.setattr(routes, "get_artifact", inaccessible)
    with pytest.raises(HTTPException) as error:
        await routes.download_artifact("missing", "mp4")
    assert error.value.status_code == 404
    assert error.value.detail == "Artifact not found"
